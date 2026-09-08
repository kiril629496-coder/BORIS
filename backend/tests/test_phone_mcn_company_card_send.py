import importlib.util
import json
import tempfile
import threading
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from docx import Document

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "phone-mcn-company-card-send.py"
SPEC = importlib.util.spec_from_file_location("phone_mcn_company_card_send", SCRIPT)
M = importlib.util.module_from_spec(SPEC)
assert SPEC and SPEC.loader
SPEC.loader.exec_module(M)


def payload():
    return {
        "org_name": "ИП Остапенко Кирилл Олегович",
        "inn": "911002277804",
        "ogrnip": "322784700115639",
        "address": "TEST ADDRESS",
        "bank_account": "11112222333344445555",
        "bank_bik": "123456789",
        "bank_name": "TEST BANK",
        "corr_account": "30101810100000000000",
        "email": "owner@example.test",
        "phone": "+70000000000",
        "updated_at": "2026-07-25T05:56:47.209180",
    }


class PhoneMcnCompanyCardSendTests(unittest.TestCase):
    def setUp(self):
        base = payload()
        self._draft_state_patcher = patch.object(
            M, "load_draft_state",
            return_value={
                "action_id": M.DRAFT_ACTION_ID,
                "mailbox_id": M.DEFAULT_MAILBOX_ID,
                "in_reply_to": "<mcn@test>",
                "card_fingerprint": M.requisites_fingerprint(base),
            },
        )
        self._draft_state_patcher.start()
        self.addCleanup(self._draft_state_patcher.stop)
        self._draft_live_patcher = patch.object(
            M, "company_card_draft_exists",
            return_value={"exists": True, "reason": "ok"},
        )
        self._draft_live_patcher.start()
        self.addCleanup(self._draft_live_patcher.stop)
        self._sent_copy_patcher = patch.object(
            M, "sent_company_card_copy_exists",
            return_value={"exists": False, "reason": "unit_test_no_sent_copy"},
        )
        self.sent_copy_guard = self._sent_copy_patcher.start()
        self.addCleanup(self._sent_copy_patcher.stop)

    def test_validate_complete_and_incomplete(self):
        p = payload()
        self.assertTrue(M.validate_requisites(p)["ok"])
        del p["bank_bik"]
        out = M.validate_requisites(p)
        self.assertFalse(out["ok"])
        self.assertIn("bank_bik", out["missing"])

    def test_requisites_fingerprint_is_deterministic_and_changes_with_data(self):
        p = payload()
        first = M.requisites_fingerprint(p)
        self.assertEqual(first, M.requisites_fingerprint(dict(p)))
        changed = dict(p)
        changed["bank_name"] = "OTHER TEST BANK"
        self.assertNotEqual(first, M.requisites_fingerprint(changed))
        self.assertEqual(len(first), 64)

    def test_build_company_card_contains_all_labels(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "card.docx"
            M.build_company_card(payload(), path)
            self.assertTrue(path.is_file())
            doc = Document(path)
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [c.text for t in doc.tables for row in t.rows for c in row.cells]
            )
            for label in [
                "Полное наименование", "ИНН", "ОГРНИП", "Адрес",
                "Расчетный счет", "Банк", "БИК",
                "Корреспондентский счет", "Email", "Телефон",
            ]:
                self.assertIn(label, text)
            self.assertIn("ИП Остапенко Кирилл Олегович", text)

    def test_safe_summary_never_contains_banking_values(self):
        p = payload()
        out = M.safe_summary(p, M.validate_requisites(p))
        raw = json.dumps(out, ensure_ascii=False)
        self.assertNotIn(p["bank_account"], raw)
        self.assertNotIn(p["bank_bik"], raw)
        self.assertNotIn(p["corr_account"], raw)
        self.assertTrue(out["banking_present"])

    def test_send_requires_explicit_banking_confirmation(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch("sys.argv", ["prog", "--apply"]):
            rc = M.main()
        self.assertEqual(rc, 4)

    def test_apply_sends_attachment_only_after_confirmation(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "claim_send_once", return_value={"claimed": True, "status": "sending"}), \
             patch.object(M, "finish_send_state") as finish, \
             patch.object(M, "send_outbound", return_value=(True, "ok", "<sent@test>")) as send, \
             patch.object(M, "sent_copy_saved", return_value=True), \
             patch.object(M, "remove_current_company_card_draft", return_value={"removed":True,"reason":"ok","count":1}) as cleanup, \
             patch("sys.argv", ["prog", "--apply", "--confirm-share-banking"]):
            rc = M.main()
        self.assertEqual(rc, 0)
        kwargs = send.call_args.kwargs
        self.assertEqual(len(kwargs["attachments"]), 1)
        self.assertEqual(kwargs["headers"]["In-Reply-To"], "<mcn@test>")
        self.assertEqual(finish.call_args.kwargs["status"], "accepted")
        cleanup.assert_called_once_with(M.DEFAULT_MAILBOX_ID)

    def test_invalid_banking_format_fails_closed(self):
        p = payload()
        p["bank_bik"] = "12"
        out = M.validate_requisites(p)
        self.assertFalse(out["ok"])
        self.assertIn("bank_bik", out["invalid_format"])

    def test_non_mcn_recipient_is_rejected_before_mail_access(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers") as thread, \
             patch("sys.argv", ["prog", "--recipient", "outside@example.test", "--apply", "--confirm-share-banking"]):
            rc = M.main()
        self.assertEqual(rc, 6)
        thread.assert_not_called()


    def test_save_draft_binds_exact_guardian_request_to_thread_lookup(self):
        p = payload()
        request_mid = "<exact-request@mcn.ru>"
        request_date = "Tue, 8 Sep 2026 11:59:05 +0300"
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": request_mid,
                 "References": request_mid,
                 "_request_date": request_date,
             }) as thread, \
             patch.object(M, "company_card_draft_exists", return_value={"exists": False, "reason": "ok"}), \
             patch.object(M, "save_draft", return_value=(True, "ok", "<draft@test>")), \
             patch.object(M, "persist_draft_state"), \
             patch.object(M, "remove_legacy_company_card_draft", return_value={
                 "removed": False, "reason": "legacy_draft_not_found", "count": 0,
             }), \
             patch("sys.argv", [
                 "prog", "--save-draft",
                 "--request-message-id", request_mid,
                 "--request-date", request_date,
             ]):
            rc = M.main()
        self.assertEqual(rc, 0)
        thread.assert_called_once_with(
            M.DEFAULT_MAILBOX_ID,
            request_message_id=request_mid,
            request_date=request_date,
        )

    def test_save_draft_never_calls_smtp_send(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "company_card_draft_exists", return_value={"exists": False, "reason": "ok"}), \
             patch.object(M, "save_draft", return_value=(True, "ok", "<draft@test>")) as draft, \
             patch.object(M, "persist_draft_state") as persist, \
             patch.object(M, "remove_legacy_company_card_draft", return_value={"removed":True,"reason":"ok","count":1}) as cleanup, \
             patch.object(M, "send_outbound") as send, \
             patch("sys.argv", ["prog", "--save-draft"]):
            rc = M.main()
        self.assertEqual(rc, 0)
        draft.assert_called_once()
        kwargs = draft.call_args.kwargs
        self.assertEqual(kwargs["attachments"][0]["filename"], "Карточка_ИП_актуальная.docx")
        cleanup.assert_called_once_with(M.DEFAULT_MAILBOX_ID)
        send.assert_not_called()

    def test_save_draft_is_idempotent_when_existing_draft_is_fresh(self):
        p = payload()
        fingerprint = M.requisites_fingerprint(p)
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "company_card_draft_exists", return_value={
                 "exists": True, "reason": "ok", "source": "persisted_message_id",
                 "card_fingerprint": fingerprint,
             }), \
             patch.object(M, "remove_current_company_card_draft") as cleanup, \
             patch.object(M, "save_draft") as draft, \
             patch("sys.argv", ["prog", "--save-draft"]):
            rc = M.main()
        self.assertEqual(rc, 0)
        cleanup.assert_not_called()
        draft.assert_not_called()

    def test_manual_sent_copy_blocks_new_draft_creation(self):
        p = payload()
        self.sent_copy_guard.return_value = {
            "exists": True,
            "reason": "matching_boris_sent_copy",
        }
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "remove_current_company_card_draft", return_value={
                 "removed": True, "reason": "ok", "count": 1,
             }) as cleanup, \
             patch.object(M, "save_draft") as draft, \
             patch("sys.argv", ["prog", "--save-draft"]):
            rc = M.main()
        self.assertEqual(rc, 0)
        cleanup.assert_called_once_with(M.DEFAULT_MAILBOX_ID)
        draft.assert_not_called()

    def test_concurrent_draft_creation_serializes_external_append(self):
        fingerprint = M.requisites_fingerprint(payload())
        shared = {"exists": False, "save_calls": 0}
        state_lock = threading.Lock()

        def fake_exists(_mailbox_id, _in_reply_to=""):
            with state_lock:
                if shared["exists"]:
                    return {
                        "exists": True,
                        "reason": "ok",
                        "source": "persisted_message_id",
                        "card_fingerprint": fingerprint,
                    }
                return {"exists": False, "reason": "ok"}

        def fake_save(*_args, **_kwargs):
            with state_lock:
                shared["save_calls"] += 1
                call_no = shared["save_calls"]
            # Keep the first worker inside the external boundary long enough
            # for the second worker to contend on the PostgreSQL advisory lock.
            time.sleep(0.15)
            return True, "ok", f"<draft-{call_no}@test>"

        def fake_persist(_mailbox_id, _message_id, _reply_id, _fingerprint):
            with state_lock:
                shared["exists"] = True

        args = (
            M.DEFAULT_MAILBOX_ID,
            M.DEFAULT_RECIPIENT,
            {"In-Reply-To": "<mcn-concurrency-qa@test>", "References": "<mcn-concurrency-qa@test>"},
            "Re: MCN QA",
            "body",
            [],
        )

        with patch.object(M, "send_state_for_thread", return_value={"status": "none", "matches": False}), \
             patch.object(M, "sent_company_card_copy_exists", return_value={"exists": False, "reason": "none"}), \
             patch.object(M, "company_card_draft_exists", side_effect=fake_exists), \
             patch.object(M, "save_draft", side_effect=fake_save), \
             patch.object(M, "persist_draft_state", side_effect=fake_persist), \
             patch.object(M, "remove_legacy_company_card_draft", return_value={
                 "removed": False, "reason": "legacy_draft_not_found", "count": 0,
             }):
            def run_one():
                summary = {}
                return M.prepare_company_card_draft_once(
                    *args,
                    summary,
                    fingerprint,
                )

            with ThreadPoolExecutor(max_workers=2) as pool:
                results = list(pool.map(lambda _x: run_one(), range(2)))

        self.assertEqual(results, [0, 0])
        self.assertEqual(shared["save_calls"], 1)


    def test_save_draft_replaces_stale_requisites_fingerprint(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "company_card_draft_exists", return_value={
                 "exists": True, "reason": "ok", "source": "persisted_message_id",
                 "card_fingerprint": "stale-fingerprint",
             }), \
             patch.object(M, "remove_current_company_card_draft", return_value={
                 "removed": True, "reason": "ok", "count": 1,
             }) as cleanup, \
             patch.object(M, "save_draft", return_value=(True, "ok", "<new-draft@test>")) as draft, \
             patch.object(M, "persist_draft_state") as persist, \
             patch.object(M, "remove_legacy_company_card_draft", return_value={
                 "removed": False, "reason": "legacy_draft_not_found", "count": 0,
             }), \
             patch("sys.argv", ["prog", "--save-draft"]):
            rc = M.main()
        self.assertEqual(rc, 0)
        cleanup.assert_called_once_with(M.DEFAULT_MAILBOX_ID)
        draft.assert_called_once()
        persist.assert_called_once()

    def test_save_draft_fails_closed_if_stale_persisted_draft_cannot_be_removed(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "company_card_draft_exists", return_value={
                 "exists": True, "reason": "ok", "source": "persisted_message_id",
                 "card_fingerprint": "stale-fingerprint",
             }), \
             patch.object(M, "remove_current_company_card_draft", return_value={
                 "removed": False, "reason": "imap_error", "count": 0,
             }), \
             patch.object(M, "save_draft") as draft, \
             patch("sys.argv", ["prog", "--save-draft"]):
            rc = M.main()
        self.assertEqual(rc, 9)
        draft.assert_not_called()


    def test_apply_blocks_when_prepared_card_fingerprint_is_stale(self):
        p = payload()
        stale = dict(p)
        stale["bank_name"] = "OLD BANK"
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "load_draft_state", return_value={
                 "action_id": M.DRAFT_ACTION_ID,
                 "mailbox_id": M.DEFAULT_MAILBOX_ID,
                 "in_reply_to": "<mcn@test>",
                 "card_fingerprint": M.requisites_fingerprint(stale),
             }), \
             patch.object(M, "claim_send_once") as claim, \
             patch.object(M, "send_outbound") as send, \
             patch("sys.argv", ["prog", "--apply", "--confirm-share-banking"]):
            rc = M.main()
        self.assertEqual(rc, 9)
        claim.assert_not_called()
        send.assert_not_called()

    def test_apply_blocks_when_prepared_draft_is_missing(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "company_card_draft_exists", return_value={"exists": False}), \
             patch.object(M, "claim_send_once") as claim, \
             patch.object(M, "send_outbound") as send, \
             patch("sys.argv", ["prog", "--apply", "--confirm-share-banking"]):
            rc = M.main()
        self.assertEqual(rc, 9)
        claim.assert_not_called()
        send.assert_not_called()

    def test_smtp_accepted_without_sent_copy_never_requests_resend(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "claim_send_once", return_value={"claimed": True, "status": "sending"}), \
             patch.object(M, "finish_send_state") as finish, \
             patch.object(M, "send_outbound", return_value=(True, "ok_sent_copy_pending", "<sent@test>")), \
             patch.object(M, "sent_copy_saved", return_value=False), \
             patch.object(M, "remove_current_company_card_draft", return_value={"removed":True,"reason":"ok","count":1}), \
             patch("sys.argv", ["prog", "--apply", "--confirm-share-banking"]):
            rc = M.main()
        self.assertEqual(rc, 0)
        self.assertEqual(finish.call_args.kwargs["status"], "accepted")

    def test_existing_accepted_claim_blocks_duplicate_smtp_send(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "claim_send_once", return_value={
                 "claimed": False, "status": "accepted", "message_id": "<sent@test>"
             }), \
             patch.object(M, "send_outbound") as send, \
             patch.object(M, "sent_copy_saved", return_value=True), \
             patch("sys.argv", ["prog", "--apply", "--confirm-share-banking"]):
            rc = M.main()
        self.assertEqual(rc, 0)
        send.assert_not_called()

    def test_surviving_sending_claim_blocks_duplicate_smtp_send(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "claim_send_once", return_value={
                 "claimed": False, "status": "sending"
             }), \
             patch.object(M, "send_outbound") as send, \
             patch("sys.argv", ["prog", "--apply", "--confirm-share-banking"]):
            rc = M.main()
        self.assertEqual(rc, 8)
        send.assert_not_called()

    def test_delivery_unknown_becomes_ambiguous_and_blocks_retry(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "claim_send_once", return_value={"claimed": True, "status": "sending"}), \
             patch.object(M, "finish_send_state") as finish, \
             patch.object(M, "send_outbound", return_value=(False, "delivery_unknown:TimeoutError", "<maybe@test>")), \
             patch("sys.argv", ["prog", "--apply", "--confirm-share-banking"]):
            rc = M.main()
        self.assertEqual(rc, 8)
        self.assertEqual(finish.call_args.kwargs["status"], "ambiguous")

    def test_proven_pre_data_failure_is_retryable(self):
        p = payload()
        with patch.object(M, "load_owner_requisites", return_value=p), \
             patch.object(M, "latest_mcn_thread_headers", return_value={
                 "In-Reply-To": "<mcn@test>", "References": "<mcn@test>"
             }), \
             patch.object(M, "claim_send_once", return_value={"claimed": True, "status": "sending"}), \
             patch.object(M, "finish_send_state") as finish, \
             patch.object(M, "send_outbound", return_value=(False, "SMTPRecipientsRefused", "<failed@test>")), \
             patch("sys.argv", ["prog", "--apply", "--confirm-share-banking"]):
            rc = M.main()
        self.assertEqual(rc, 5)
        self.assertEqual(finish.call_args.kwargs["status"], "safe_failure")


class PhoneMcnCompanyCardSentEvidenceTests(unittest.TestCase):
    def _mailbox_row(self):
        return {
            "secret_encrypted": "enc",
            "imap_ssl": True,
            "imap_host": "imap.invalid",
            "imap_port": 993,
            "username": "owner@example.test",
        }

    def _cp(self, items):
        return type("CP", (), {
            "returncode": 0,
            "stdout": json.dumps({"status": "ok", "items": items}),
            "stderr": "",
        })()

    def test_discovery_fallback_rejects_matching_card_from_wrong_thread(self):
        items = [{
            "date": "Mon, 7 Sep 2026 17:00:00 +0300",
            "company_card_signal": True,
            "current_card_match": True,
            "in_reply_to": "<other-request@mcn.ru>",
        }]
        with patch.object(M, "_row", return_value=self._mailbox_row()), \
             patch.object(M, "decrypt_secret", return_value="secret"), \
             patch.object(M.imaplib, "IMAP4_SSL", side_effect=RuntimeError("imap direct unavailable")), \
             patch.object(M.subprocess, "run", return_value=self._cp(items)):
            out = M.sent_company_card_copy_exists(
                2,
                "<request@mcn.ru>",
                "a" * 64,
                "Mon, 7 Sep 2026 16:07:30 +0300",
            )
        self.assertFalse(out["exists"])
        self.assertEqual(out["reason"], "matching_sent_copy_not_found")

    def test_discovery_fallback_accepts_matching_card_only_for_exact_thread(self):
        items = [{
            "date": "Mon, 7 Sep 2026 17:00:00 +0300",
            "company_card_signal": True,
            "current_card_match": True,
            "in_reply_to": "<request@mcn.ru>",
        }]
        with patch.object(M, "_row", return_value=self._mailbox_row()), \
             patch.object(M, "decrypt_secret", return_value="secret"), \
             patch.object(M.imaplib, "IMAP4_SSL", side_effect=RuntimeError("imap direct unavailable")), \
             patch.object(M.subprocess, "run", return_value=self._cp(items)):
            out = M.sent_company_card_copy_exists(
                2,
                "<request@mcn.ru>",
                "a" * 64,
                "Mon, 7 Sep 2026 16:07:30 +0300",
            )
        self.assertTrue(out["exists"])
        self.assertEqual(out["reason"], "attachment_seen_after_request")


if __name__ == "__main__":
    unittest.main()
