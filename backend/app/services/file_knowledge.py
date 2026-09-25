# -*- coding: utf-8 -*-
"""Извлечение знаний из файлов клиента.

Клиент загружает прайс, описание услуг, договор — текст вынимается и режется
на смысловые куски, каждый становится фактом в базе знаний.

Разбор БЕЗ модели: нарезка механическая, по абзацам. Грубее, чем разбор
моделью, но работает бесплатно и не зависит от доступности OpenAI. Когда
модель доступна, поверх этого можно добавить нормальное структурирование.
"""
import re

MAX_CHUNK = 700          # символов в одном факте
MIN_CHUNK = 40           # короче — не факт, а обрывок
MAX_FACTS_PER_FILE = 120  # предохранитель от договора на 200 страниц

_WS = re.compile(r"[ \t\u00a0]+")
_MULTINL = re.compile(r"\n{2,}")


def extract_text(filename, data):
    """Текст из файла. Возвращает (текст, тип). Тип нужен для честного отказа."""
    low = (filename or "").lower()
    if low.endswith(".pdf"):
        import io
        import pdfplumber
        out = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:60]:
                out.append(page.extract_text() or "")
                # таблицы важнее всего: в прайсах цены живут именно в них
                for tbl in (page.extract_tables() or []):
                    for row in tbl:
                        cells = [str(c).strip() for c in row if c and str(c).strip()]
                        if len(cells) >= 2:
                            out.append(" — ".join(cells))
        return "\n".join(out), "pdf"
    if low.endswith(".docx"):
        import io
        import docx
        doc = docx.Document(io.BytesIO(data))
        out = [p.text for p in doc.paragraphs if p.text.strip()]
        for tbl in doc.tables:
            for row in tbl.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if len(cells) >= 2:
                    out.append(" — ".join(cells))
        return "\n".join(out), "docx"
    if low.endswith(".xlsx") or low.endswith(".xlsm"):
        import io
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        out = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
                if len(cells) >= 2:
                    out.append(" — ".join(cells))
        return "\n".join(out), "xlsx"
    if low.endswith(".txt") or low.endswith(".md") or low.endswith(".csv"):
        for enc in ("utf-8", "cp1251"):
            try:
                return data.decode(enc), "text"
            except Exception:
                continue
        return "", "text"
    return "", "unsupported"


def _category_of(block):
    """Грубая разметка без модели — по словам-приметам."""
    low = block.lower()
    if re.search(r"\d[\d\s]{2,}\s*(?:₽|руб|р\.)", low) or "цена" in low or "стоимость" in low:
        return "cena"
    if "гаранти" in low:
        return "garantiya"
    if "доставк" in low or "оплат" in low or "монтаж" in low or "услуг" in low:
        return "usluga"
    if "преимуществ" in low or "почему мы" in low:
        return "preimushchestvo"
    return "faq"


_CHAR_FIX = ((u"\u0138", u"\u043a"), (u"\ufb01", u"fi"), (u"\ufb02", u"fl"),
             (u"\u00ad", u""), (u"\r", u""))


def _fix_chars(text):
    """PDF отдаёт похожие по виду, но чужие символы: вместо «к» приходит ĸ.
    Такие факты потом не находятся поиском — чиним до разбора."""
    text = text or ""
    for _bad, _good in _CHAR_FIX:
        text = text.replace(_bad, _good)
    return text


def _name_of(block):
    """Название факта — первая осмысленная строка, коротко."""
    _lines = [l.strip() for l in block.strip().split("\n") if l.strip()]
    first = ""
    for _l in _lines[:4]:
        if len(_l) >= 12:
            first = _l
            break
    if not first:
        first = _lines[0] if _lines else ""
    first = _WS.sub(" ", first).strip(" .;:—-")
    if len(first) > 70:
        first = first[:67].rsplit(" ", 1)[0] + "…"
    return first or "Из файла"


def to_facts(text, filename):
    """Текст → список фактов. Каждый: category, name, value, snippet."""
    text = _fix_chars(text)
    blocks, buf = [], []
    for para in _MULTINL.split(text):
        for line in para.split("\n"):
            line = _WS.sub(" ", line).strip()
            if not line:
                continue
            buf.append(line)
            if sum(len(x) for x in buf) >= MAX_CHUNK:
                blocks.append("\n".join(buf))
                buf = []
        if buf:
            blocks.append("\n".join(buf))
            buf = []
    if buf:
        blocks.append("\n".join(buf))

    facts = []
    for b in blocks:
        b = b.strip()
        if len(b) < MIN_CHUNK:
            continue
        facts.append({
            "category": _category_of(b),
            "name": _name_of(b),
            "value": b[:2000],
            "snippet": ("%s: %s" % (filename, b[:180])),
        })
        if len(facts) >= MAX_FACTS_PER_FILE:
            break
    return facts
