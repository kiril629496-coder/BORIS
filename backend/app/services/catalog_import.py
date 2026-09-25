"""Импорт ассортимента в каталог ЕЦО: XLSX и CSV.

Существующий import_service (43 строки, list[str] без файлов, маппинга и
владельца) для этого не годится, но и второй очереди мы не заводим: тяжёлая
часть идёт обработчиком в тех же background_jobs.

Правило Кирилла: сначала ПРЕВЬЮ маппинга колонок и первые 20 строк, правка
человеком, и только потом запись. Сомнительные колонки молча не импортируются.
"""
import csv
import io
import re

# Синонимы колонок. Низкая уверенность = колонка не подставляется молча.
COLUMN_HINTS = {
    "name": ("название", "наименование", "товар", "модель", "name", "title"),
    "sku": ("артикул", "sku", "код", "id"),
    "price": ("цена", "стоимость", "price", "розница"),
    "price_from": ("цена от", "от", "price from", "минимальная"),
    "price_to": ("цена до", "до", "price to", "максимальная"),
    "category": ("категория", "группа", "раздел", "тип", "category"),
    "description": ("описание", "текст", "description"),
    "geo": ("город", "гео", "регион", "city"),
}
NUMERIC = ("price", "price_from", "price_to")


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "").strip().lower())


def guess_mapping(headers):
    """Колонка файла -> поле товара + уверенность. Ничего не выдумывает."""
    out = []
    for i, h in enumerate(headers):
        low = _norm(h)
        field, conf = None, 0.0
        for f, hints in COLUMN_HINTS.items():
            for hint in hints:
                if low == hint:
                    field, conf = f, 1.0
                    break
                if hint in low and conf < 0.6:
                    field, conf = f, 0.6
            if conf == 1.0:
                break
        out.append({"index": i, "header": str(h or ""), "field": field,
                    "confidence": conf,
                    "needs_confirm": bool(field) and conf < 1.0})
    return out


def read_table(content: bytes, filename: str, limit=0):
    """(headers, rows). XLSX через openpyxl, CSV с автоопределением разделителя."""
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = []
        headers = []
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n == 0:
                headers = [("" if c is None else str(c)) for c in row]
                continue
            if row is None or all(c is None for c in row):
                continue
            rows.append([("" if c is None else str(c)) for c in row])
            if limit and len(rows) >= limit:
                break
        wb.close()
        return headers, rows
    text = content.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except Exception:
        class dialect:      # noqa: N801 — запасной разделитель
            delimiter = ";"
    reader = csv.reader(io.StringIO(text), delimiter=dialect.delimiter)
    rows = list(reader)
    if not rows:
        return [], []
    headers, body = rows[0], [r for r in rows[1:] if any(str(c).strip() for c in r)]
    return headers, (body[:limit] if limit else body)


def _num(v):
    s = re.sub(r"[^\d.,]", "", str(v or "")).replace(",", ".")
    try:
        return float(s) if s else None
    except ValueError:
        return None


def preview(content: bytes, filename: str, sample=20):
    headers, rows = read_table(content, filename, limit=sample)
    mapping = guess_mapping(headers)
    known = {m["field"] for m in mapping if m["field"]}
    problems = []
    if "name" not in known:
        problems.append("не нашёл колонку с названием — укажите её вручную")
    if not rows:
        problems.append("в файле нет строк с данными")
    return {"headers": headers, "mapping": mapping,
            "rows": rows[:sample], "sample_count": len(rows),
            "problems": problems}


def apply_import(db, owner_user_id, account_id, content, filename, mapping,
                 created_by=None, dry_run=True, limit=0):
    """Пишет товары по ПОДТВЕРЖДЁННОМУ маппингу. Повтор не плодит дубли."""
    from app.services import catalog_service as C
    from app.models.catalog_product import CatalogProduct

    headers, rows = read_table(content, filename, limit=limit)
    by_field = {}
    for m in (mapping or []):
        if m.get("field"):
            by_field[m["field"]] = int(m["index"])
    if "name" not in by_field:
        return {"error": "не указана колонка с названием", "created": 0}

    created = skipped = updated = 0
    problems = []
    for n, row in enumerate(rows):
        def val(field):
            i = by_field.get(field)
            return row[i] if i is not None and i < len(row) else None
        name = str(val("name") or "").strip()
        if not name:
            skipped += 1
            continue
        sku = str(val("sku") or "").strip() or None
        # Идемпотентность: тот же владелец + sku (или имя) = тот же товар.
        q = db.query(CatalogProduct).filter(
            CatalogProduct.owner_user_id == owner_user_id,
            CatalogProduct.deleted_at.is_(None))
        exists = (q.filter(CatalogProduct.sku == sku).first() if sku
                  else q.filter(CatalogProduct.name == name).first())
        if exists is not None:
            updated += 1
            continue
        if dry_run:
            created += 1
            continue
        try:
            with db.begin_nested():
                C.create_product(
                    db, owner_user_id, account_id, name, created_by=created_by,
                    sku=sku, category=str(val("category") or "").strip() or None,
                    description=str(val("description") or "").strip() or None,
                    price=_num(val("price")), price_from=_num(val("price_from")),
                    price_to=_num(val("price_to")),
                    geo=[g.strip() for g in str(val("geo") or "").split(",") if g.strip()] or None,
                    source="import", source_ref=filename, status="draft")
            created += 1
        except Exception as e:
            skipped += 1
            if len(problems) < 5:
                problems.append("строка %d: %s" % (n + 2, repr(e)[:120]))
    if not dry_run:
        db.commit()
    return {"created": created, "updated_existing": updated, "skipped": skipped,
            "total_rows": len(rows), "problems": problems, "dry_run": dry_run}
