#!/root/BORIS/backend/venv/bin/python
from __future__ import annotations

import argparse
import email
import hashlib
import hmac
import imaplib
import json
import os
import subprocess
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from email.utils import parseaddr, parsedate_to_datetime

from dotenv import load_dotenv
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from sqlalchemy import text

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env", override=False)

from app.db.session import SessionLocal
from app.services.client_mailboxes import _drafts_folder, _sent_folder, _row, decrypt_secret, save_draft, send_outbound, sent_copy_saved

REQUIRED = (
    "org_name", "inn", "ogrnip", "address",
    "bank_account", "bank_bik", "bank_name", "corr_account",
    "email", "phone",
)
DEFAULT_RECIPIENT = "ava@mcn.ru"
DEFAULT_MAILBOX_ID = 2
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DRAFT_ACTION_ID = "boris-mcn-company-card-v4"
DRAFT_STATE_ACCOUNT = "__mcn_mailbox_watch__"
DRAFT_STATE_KEY = "mcn_company_card_draft_v4"
LEGACY_DRAFT_STATES = (
    ("mcn_company_card_draft_v1", "boris-mcn-company-card-v1"),
    ("mcn_company_card_draft_v2", "boris-mcn-company-card-v2"),
    ("mcn_company_card_draft_v3", "boris-mcn-company-card-v3"),
)
SEND_STATE_KEY = "mcn_company_card_send_v1"


def load_owner_requisites() -> dict[str, Any]:
    db = SessionLocal()
    try:
        row = db.execute(text("""
            SELECT value FROM storage
            WHERE account_id='__owner' AND key='requisites'
            ORDER BY id DESC LIMIT 1
        """)).first()
    finally:
        db.close()
    if not row or not row[0]:
        raise RuntimeError("owner_requisites_missing")
    payload = json.loads(str(row[0]))
    if not isinstance(payload, dict):
        raise RuntimeError("owner_requisites_invalid")
    return payload


def requisites_fingerprint(payload: dict[str, Any]) -> str:
    canonical = {
        key: str(payload.get(key) or "").strip()
        for key in REQUIRED
    }
    canonical["updated_at"] = str(payload.get("updated_at") or "").strip()
    raw = json.dumps(canonical, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def validate_requisites(payload: dict[str, Any]) -> dict[str, Any]:
    missing = [k for k in REQUIRED if not str(payload.get(k) or "").strip()]
    digits = lambda key: "".join(ch for ch in str(payload.get(key) or "") if ch.isdigit())
    invalid_format: list[str] = []
    if "inn" not in missing and len(digits("inn")) not in {10, 12}:
        invalid_format.append("inn")
    if "ogrnip" not in missing and len(digits("ogrnip")) != 15:
        invalid_format.append("ogrnip")
    if "bank_account" not in missing and len(digits("bank_account")) != 20:
        invalid_format.append("bank_account")
    if "bank_bik" not in missing and len(digits("bank_bik")) != 9:
        invalid_format.append("bank_bik")
    if "corr_account" not in missing and len(digits("corr_account")) != 20:
        invalid_format.append("corr_account")
    email_value = str(payload.get("email") or "").strip()
    if "email" not in missing and ("@" not in email_value or "." not in email_value.rsplit("@", 1)[-1]):
        invalid_format.append("email")
    phone_digits = digits("phone")
    if "phone" not in missing and not (10 <= len(phone_digits) <= 15):
        invalid_format.append("phone")
    return {
        "ok": not missing and not invalid_format,
        "missing": missing,
        "invalid_format": invalid_format,
        "fields_present": len(REQUIRED) - len(missing),
        "fields_required": len(REQUIRED),
        "updated_at": str(payload.get("updated_at") or "")[:80] or None,
    }


def validate_mcn_recipient(value: str) -> dict[str, Any]:
    address = parseaddr(str(value or ""))[1].strip().lower()
    domain = address.rsplit("@", 1)[-1] if "@" in address else ""
    ok = bool(address and (domain == "mcn.ru" or domain.endswith(".mcn.ru")))
    return {"ok": ok, "address": address if ok else None, "domain": domain if ok else None}


def build_company_card(payload: dict[str, Any], output: Path) -> Path:
    check = validate_requisites(payload)
    if not check["ok"]:
        raise RuntimeError("owner_requisites_incomplete:" + ",".join(check["missing"]))

    doc = Document()
    cp = doc.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.title = "Карточка индивидуального предпринимателя"
    cp.subject = "Реквизиты для оформления услуг MCN"

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Карточка индивидуального предпринимателя")
    run.bold = True
    run.font.size = Pt(15)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(str(payload["org_name"]).strip()).bold = True

    fields = [
        ("Полное наименование", "org_name"),
        ("ИНН", "inn"),
        ("ОГРНИП", "ogrnip"),
        ("Адрес", "address"),
        ("Расчетный счет", "bank_account"),
        ("Банк", "bank_name"),
        ("БИК", "bank_bik"),
        ("Корреспондентский счет", "corr_account"),
        ("Email", "email"),
        ("Телефон", "phone"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, key in fields:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(payload[key]).strip()
        cells[0].paragraphs[0].runs[0].bold = True

    updated = str(payload.get("updated_at") or "").strip()
    if updated:
        p = doc.add_paragraph()
        p.add_run("Актуальность реквизитов в BORIS: ").bold = True
        p.add_run(updated)

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output))
    return output


def latest_mcn_thread_headers(
    mailbox_id: int,
    request_message_id: str = "",
    request_date: str = "",
) -> dict[str, str]:
    """Return headers for one exact MCN request, or the latest request if omitted.

    MCN_EXACT_REQUEST_BINDING_V1 prevents guardian/sender races when multiple
    provider emails arrive within the same polling window. When a Message-ID is
    supplied we fail closed unless that exact message exists in the MCN mailbox.
    """
    row = _row(mailbox_id)
    if not row:
        raise RuntimeError("mailbox_not_found")
    secret = decrypt_secret(row["secret_encrypted"])
    cls = imaplib.IMAP4_SSL if row["imap_ssl"] else imaplib.IMAP4
    m = cls(row["imap_host"], int(row["imap_port"]), timeout=15)
    exact_mid = str(request_message_id or "").strip()
    expected_date = str(request_date or "").strip()
    try:
        m.login(row["username"], secret)
        st, _ = m.select("INBOX", readonly=True)
        if st != "OK":
            raise RuntimeError("imap_select_failed")
        st, data = m.uid("search", None, "FROM", '"ava@mcn.ru"')
        uids = data[0].split() if st == "OK" and data and data[0] else []
        if not uids:
            raise RuntimeError("mcn_thread_not_found")

        ordered = sorted(uids, key=lambda x: int(x), reverse=True)
        # An explicitly selected operational request is always recent; keep the
        # read bounded while still allowing enough history for delayed retries.
        candidates = ordered[:250] if exact_mid else ordered[:1]
        for uid in candidates:
            st, fetched = m.uid(
                "fetch", uid,
                "(BODY.PEEK[HEADER.FIELDS (DATE MESSAGE-ID REFERENCES)])",
            )
            raw = b"".join(
                x[1] for x in fetched
                if isinstance(x, tuple) and len(x) > 1 and isinstance(x[1], (bytes, bytearray))
            ) if st == "OK" else b""
            if not raw:
                continue
            msg = email.message_from_bytes(raw)
            message_id = str(msg.get("Message-ID") or "").strip()
            if exact_mid and message_id != exact_mid:
                continue
            references = str(msg.get("References") or "").strip()
            actual_date = str(msg.get("Date") or "").strip()
            if not message_id:
                continue
            refs = " ".join(x for x in [references, message_id] if x).strip()
            return {
                "In-Reply-To": message_id,
                "References": refs,
                "_request_date": actual_date or expected_date,
            }

        if exact_mid:
            raise RuntimeError("mcn_exact_request_not_found")
        raise RuntimeError("mcn_message_id_missing")
    finally:
        try:
            m.logout()
        except Exception:
            pass


def load_draft_state(state_key: str = DRAFT_STATE_KEY) -> dict[str, Any]:
    db = SessionLocal()
    try:
        row = db.execute(text("""
            SELECT value FROM storage
            WHERE account_id=:a AND key=:k
            ORDER BY id DESC LIMIT 1
        """), {"a": DRAFT_STATE_ACCOUNT, "k": str(state_key)}).first()
    finally:
        db.close()
    if not row or not row[0]:
        return {}
    try:
        payload = json.loads(str(row[0]))
    except Exception:
        return {}
    return payload if isinstance(payload, dict) else {}


def persist_draft_state(mailbox_id: int, message_id: str, in_reply_to: str, card_fingerprint: str) -> None:
    payload = {
        "action_id": DRAFT_ACTION_ID,
        "mailbox_id": int(mailbox_id),
        "message_id": str(message_id or "")[:500],
        "in_reply_to": str(in_reply_to or "")[:500],
        "card_fingerprint": str(card_fingerprint or "")[:64],
        "saved_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
    }
    raw = json.dumps(payload, ensure_ascii=False)
    db = SessionLocal()
    try:
        row = db.execute(text("""
            SELECT id FROM storage
            WHERE account_id=:a AND key=:k
            ORDER BY id DESC LIMIT 1
        """), {"a": DRAFT_STATE_ACCOUNT, "k": DRAFT_STATE_KEY}).first()
        if row:
            db.execute(
                text("UPDATE storage SET value=:v WHERE id=:i"),
                {"v": raw, "i": int(row[0])},
            )
        else:
            db.execute(
                text("INSERT INTO storage(account_id,key,value) VALUES(:a,:k,:v)"),
                {"a": DRAFT_STATE_ACCOUNT, "k": DRAFT_STATE_KEY, "v": raw},
            )
        db.commit()
    finally:
        db.close()


def _decode_state(raw: Any) -> dict[str, Any]:
    try:
        payload = json.loads(str(raw or ""))
    except Exception:
        return {}
    return payload if isinstance(payload, dict) else {}


def load_send_state() -> dict[str, Any]:
    db = SessionLocal()
    try:
        row = db.execute(text("""
            SELECT value FROM storage
            WHERE account_id=:a AND key=:k
            ORDER BY id DESC LIMIT 1
        """), {"a": DRAFT_STATE_ACCOUNT, "k": SEND_STATE_KEY}).first()
    finally:
        db.close()
    return _decode_state(row[0]) if row else {}


def send_state_for_thread(mailbox_id: int, in_reply_to: str) -> dict[str, Any]:
    state = load_send_state()
    reply_id = str(in_reply_to or "").strip()[:500]
    same = (
        state.get("action_id") == DRAFT_ACTION_ID
        and int(state.get("mailbox_id") or 0) == int(mailbox_id)
        and str(state.get("in_reply_to") or "") == reply_id
    )
    if not same:
        return {"status": "none", "matches": False}
    return {
        "status": str(state.get("status") or "")[:40] or "none",
        "matches": True,
        "message_id": str(state.get("message_id") or "")[:500] or None,
        "reason": str(state.get("reason") or "")[:160] or None,
    }


def claim_send_once(mailbox_id: int, in_reply_to: str) -> dict[str, Any]:
    """Atomically reserve the external send boundary.

    A surviving "sending" claim is deliberately ambiguous after a crash and
    therefore blocks automatic resend. Only failures proven to happen before
    SMTP DATA acceptance are made retryable by finish_send_state(...safe_failure).
    """
    reply_id = str(in_reply_to or "").strip()[:500]
    if not reply_id:
        return {"claimed": False, "status": "invalid_thread"}
    db = SessionLocal()
    try:
        lock_key = f"mcn-company-card-send|{int(mailbox_id)}|{reply_id}"
        db.execute(text("SELECT pg_advisory_xact_lock(hashtextextended(:k,0))"), {"k": lock_key})
        row = db.execute(text("""
            SELECT id,value FROM storage
            WHERE account_id=:a AND key=:k
            ORDER BY id DESC LIMIT 1
            FOR UPDATE
        """), {"a": DRAFT_STATE_ACCOUNT, "k": SEND_STATE_KEY}).first()
        current = _decode_state(row[1]) if row else {}
        same = (
            current.get("action_id") == DRAFT_ACTION_ID
            and int(current.get("mailbox_id") or 0) == int(mailbox_id)
            and str(current.get("in_reply_to") or "") == reply_id
        )
        status = str(current.get("status") or "")
        if same and status in {"sending", "accepted", "ambiguous"}:
            db.commit()
            return {
                "claimed": False,
                "status": status,
                "message_id": str(current.get("message_id") or "")[:500] or None,
                "reason": str(current.get("reason") or "")[:160] or None,
            }
        payload = {
            "action_id": DRAFT_ACTION_ID,
            "mailbox_id": int(mailbox_id),
            "in_reply_to": reply_id,
            "status": "sending",
            "message_id": "",
            "reason": "owner_approved_send_started",
            "updated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        }
        raw = json.dumps(payload, ensure_ascii=False)
        if row:
            db.execute(text("UPDATE storage SET value=:v WHERE id=:i"), {"v": raw, "i": int(row[0])})
        else:
            db.execute(
                text("INSERT INTO storage(account_id,key,value) VALUES(:a,:k,:v)"),
                {"a": DRAFT_STATE_ACCOUNT, "k": SEND_STATE_KEY, "v": raw},
            )
        db.commit()
        return {"claimed": True, "status": "sending"}
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def finish_send_state(
    mailbox_id: int,
    in_reply_to: str,
    *,
    status: str,
    message_id: str = "",
    reason: str = "",
) -> None:
    allowed = {"accepted", "ambiguous", "safe_failure"}
    final_status = str(status or "")
    if final_status not in allowed:
        raise ValueError("invalid_send_state")
    reply_id = str(in_reply_to or "").strip()[:500]
    db = SessionLocal()
    try:
        lock_key = f"mcn-company-card-send|{int(mailbox_id)}|{reply_id}"
        db.execute(text("SELECT pg_advisory_xact_lock(hashtextextended(:k,0))"), {"k": lock_key})
        row = db.execute(text("""
            SELECT id,value FROM storage
            WHERE account_id=:a AND key=:k
            ORDER BY id DESC LIMIT 1
            FOR UPDATE
        """), {"a": DRAFT_STATE_ACCOUNT, "k": SEND_STATE_KEY}).first()
        current = _decode_state(row[1]) if row else {}
        if not row or (
            current.get("action_id") != DRAFT_ACTION_ID
            or int(current.get("mailbox_id") or 0) != int(mailbox_id)
            or str(current.get("in_reply_to") or "") != reply_id
        ):
            raise RuntimeError("send_claim_missing_or_changed")
        current.update({
            "status": final_status,
            "message_id": str(message_id or "")[:500],
            "reason": str(reason or "")[:160],
            "updated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        })
        db.execute(
            text("UPDATE storage SET value=:v WHERE id=:i"),
            {"v": json.dumps(current, ensure_ascii=False), "i": int(row[0])},
        )
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def company_card_draft_exists(mailbox_id: int, in_reply_to: str = "") -> dict[str, Any]:
    row = _row(mailbox_id)
    if not row:
        return {"exists": False, "reason": "mailbox_not_found"}
    secret = decrypt_secret(row["secret_encrypted"])
    cls = imaplib.IMAP4_SSL if row["imap_ssl"] else imaplib.IMAP4
    m = None
    try:
        m = cls(row["imap_host"], int(row["imap_port"]), timeout=15)
        m.login(row["username"], secret)
        folder = _drafts_folder(m)
        if not folder:
            return {"exists": False, "reason": "drafts_folder_not_found"}
        st, _ = m.select(folder, readonly=True)
        if st != "OK":
            return {"exists": False, "reason": "drafts_select_failed"}
        state = load_draft_state()
        state_message_id = str(state.get("message_id") or "").strip()
        if (
            state.get("action_id") == DRAFT_ACTION_ID
            and int(state.get("mailbox_id") or 0) == int(mailbox_id)
            and state_message_id
        ):
            st, data = m.uid(
                "search", None,
                "HEADER", "Message-ID", f'"{state_message_id}"',
            )
            ids = data[0].split() if st == "OK" and data and data[0] else []
            if ids:
                return {
                    "exists": True,
                    "reason": "ok",
                    "source": "persisted_message_id",
                    "card_fingerprint": str(state.get("card_fingerprint") or "")[:64] or None,
                    "in_reply_to": str(state.get("in_reply_to") or "")[:500] or None,
                }

        st, data = m.uid(
            "search", None,
            "HEADER", "X-BORIS-Action-ID", f'"{DRAFT_ACTION_ID}"',
        )
        ids = data[0].split() if st == "OK" and data and data[0] else []
        if not ids and str(in_reply_to or "").strip():
            st, data = m.uid(
                "search", None,
                "HEADER", "In-Reply-To", f'"{str(in_reply_to).strip()}"',
            )
            ids = data[0].split() if st == "OK" and data and data[0] else []
        return {"exists": bool(ids), "reason": "ok"}
    except Exception as exc:
        return {"exists": False, "reason": type(exc).__name__[:120]}
    finally:
        if m is not None:
            try:
                m.logout()
            except Exception:
                pass


def remove_legacy_company_card_draft(mailbox_id: int) -> dict[str, Any]:
    """Remove only exact obsolete BORIS-created draft versions after current draft exists."""
    states: list[tuple[str, str, str]] = []
    for state_key, expected_action in LEGACY_DRAFT_STATES:
        state = load_draft_state(state_key)
        message_id = str(state.get("message_id") or "").strip()
        if (
            state.get("action_id") == expected_action
            and int(state.get("mailbox_id") or 0) == int(mailbox_id)
            and message_id
        ):
            states.append((state_key, expected_action, message_id))
    if not states:
        return {"removed": False, "reason": "legacy_state_not_found", "count": 0}

    row = _row(mailbox_id)
    if not row:
        return {"removed": False, "reason": "mailbox_not_found", "count": 0}
    secret = decrypt_secret(row["secret_encrypted"])
    cls = imaplib.IMAP4_SSL if row["imap_ssl"] else imaplib.IMAP4
    m = None
    removed = 0
    cleaned_keys: list[str] = []
    try:
        m = cls(row["imap_host"], int(row["imap_port"]), timeout=15)
        m.login(row["username"], secret)
        folder = _drafts_folder(m)
        if not folder:
            return {"removed": False, "reason": "drafts_folder_not_found", "count": 0}
        st, _ = m.select(folder)
        if st != "OK":
            return {"removed": False, "reason": "drafts_select_failed", "count": 0}
        for state_key, _expected_action, message_id in states:
            st, data = m.uid("search", None, "HEADER", "Message-ID", f'"{message_id}"')
            ids = data[0].split() if st == "OK" and data and data[0] else []
            key_removed = 0
            for uid in ids:
                st, _ = m.uid("store", uid, "+FLAGS.SILENT", r"(\Deleted)")
                if st == "OK":
                    key_removed += 1
            if key_removed:
                removed += key_removed
                cleaned_keys.append(state_key)
        if removed:
            m.expunge()
        if cleaned_keys:
            db = SessionLocal()
            try:
                for key in cleaned_keys:
                    db.execute(
                        text("DELETE FROM storage WHERE account_id=:a AND key=:k"),
                        {"a": DRAFT_STATE_ACCOUNT, "k": key},
                    )
                db.commit()
            finally:
                db.close()
        return {
            "removed": bool(removed),
            "reason": "ok" if removed else "legacy_draft_not_found",
            "count": int(removed),
        }
    except Exception as exc:
        return {"removed": False, "reason": type(exc).__name__[:120], "count": int(removed)}
    finally:
        if m is not None:
            try:
                m.logout()
            except Exception:
                pass


def remove_current_company_card_draft(mailbox_id: int) -> dict[str, Any]:
    """Remove only BORIS's current company-card draft.

    Prefer the persisted Message-ID. If the DB state was lost, fall back to the
    private X-BORIS-Action-ID header so stale drafts can self-heal without
    touching unrelated user drafts.
    """
    state = load_draft_state(DRAFT_STATE_KEY)
    message_id = str(state.get("message_id") or "").strip()
    state_matches = (
        state.get("action_id") == DRAFT_ACTION_ID
        and int(state.get("mailbox_id") or 0) == int(mailbox_id)
    )
    row = _row(mailbox_id)
    if not row:
        return {"removed": False, "reason": "mailbox_not_found", "count": 0}
    secret = decrypt_secret(row["secret_encrypted"])
    cls = imaplib.IMAP4_SSL if row["imap_ssl"] else imaplib.IMAP4
    m = None
    try:
        m = cls(row["imap_host"], int(row["imap_port"]), timeout=15)
        m.login(row["username"], secret)
        folder = _drafts_folder(m)
        if not folder:
            return {"removed": False, "reason": "drafts_folder_not_found", "count": 0}
        st, _ = m.select(folder)
        if st != "OK":
            return {"removed": False, "reason": "drafts_select_failed", "count": 0}

        ids = []
        if state_matches and message_id:
            st, data = m.uid("search", None, "HEADER", "Message-ID", f'"{message_id}"')
            ids = data[0].split() if st == "OK" and data and data[0] else []
        if not ids:
            st, data = m.uid(
                "search", None,
                "HEADER", "X-BORIS-Action-ID", f'"{DRAFT_ACTION_ID}"',
            )
            ids = data[0].split() if st == "OK" and data and data[0] else []

        removed = 0
        for uid in ids:
            st, _ = m.uid("store", uid, "+FLAGS.SILENT", r"(\Deleted)")
            if st == "OK":
                removed += 1
        if removed:
            m.expunge()

        if removed or state_matches:
            db = SessionLocal()
            try:
                db.execute(
                    text("DELETE FROM storage WHERE account_id=:a AND key=:k"),
                    {"a": DRAFT_STATE_ACCOUNT, "k": DRAFT_STATE_KEY},
                )
                db.commit()
            finally:
                db.close()

        return {
            "removed": bool(removed),
            "reason": "ok" if removed else "current_draft_not_found",
            "count": int(removed),
        }
    except Exception as exc:
        return {"removed": False, "reason": type(exc).__name__[:120], "count": 0}
    finally:
        if m is not None:
            try:
                m.logout()
            except Exception:
                pass


def dedupe_boris_company_card_drafts(mailbox_id: int, scan_limit: int = 500, dry_run: bool = False) -> dict[str, Any]:
    """Remove only BORIS-tagged duplicate/legacy MCN company-card drafts.

    The current persisted Message-ID is preserved. Untagged drafts are never
    touched, so a user's own Mail.ru draft cannot be deleted by this self-heal.
    """
    state = load_draft_state(DRAFT_STATE_KEY)
    keep_message_id = str(state.get("message_id") or "").strip()
    known_actions = {DRAFT_ACTION_ID}
    known_actions.update(expected for _key, expected in LEGACY_DRAFT_STATES)

    row = _row(mailbox_id)
    if not row:
        return {"status": "blocked", "reason": "mailbox_not_found", "removed": 0, "kept": 0}
    secret = decrypt_secret(row["secret_encrypted"])
    cls = imaplib.IMAP4_SSL if row["imap_ssl"] else imaplib.IMAP4
    m = None
    removed = 0
    planned = 0
    kept = 0
    tagged = 0
    try:
        m = cls(row["imap_host"], int(row["imap_port"]), timeout=15)
        m.login(row["username"], secret)
        folder = _drafts_folder(m)
        if not folder:
            return {"status": "blocked", "reason": "drafts_folder_not_found", "removed": 0, "kept": 0}
        st, _ = m.select(folder)
        if st != "OK":
            return {"status": "blocked", "reason": "drafts_select_failed", "removed": 0, "kept": 0}
        st, data = m.uid("search", None, "ALL")
        uids = data[0].split() if st == "OK" and data and data[0] else []
        for uid in uids[-max(1, min(int(scan_limit), 2000)):]:
            st, fetched = m.uid(
                "fetch",
                uid,
                "(BODY.PEEK[HEADER.FIELDS (TO MESSAGE-ID X-BORIS-ACTION-ID)])",
            )
            raw = b"".join(
                x[1] for x in fetched
                if isinstance(x, tuple) and len(x) > 1 and isinstance(x[1], (bytes, bytearray))
            ) if st == "OK" else b""
            if not raw:
                continue
            msg = email.message_from_bytes(raw)
            action = str(msg.get("X-BORIS-Action-ID") or "").strip()
            if action not in known_actions:
                continue
            recipient = parseaddr(str(msg.get("To") or ""))[1].strip().lower()
            recipient_domain = recipient.rsplit("@", 1)[-1] if "@" in recipient else ""
            if recipient_domain != "mcn.ru" and not recipient_domain.endswith(".mcn.ru"):
                continue
            tagged += 1
            message_id = str(msg.get("Message-ID") or "").strip()
            if action == DRAFT_ACTION_ID and keep_message_id and message_id == keep_message_id:
                kept += 1
                continue
            if action == DRAFT_ACTION_ID and not keep_message_id:
                kept += 1
                continue
            planned += 1
            if dry_run:
                continue
            st, _ = m.uid("store", uid, "+FLAGS.SILENT", r"(\Deleted)")
            if st == "OK":
                removed += 1
        if removed:
            m.expunge()
        return {
            "status": "ok",
            "reason": "ok",
            "tagged": int(tagged),
            "planned": int(planned),
            "removed": int(removed),
            "kept": int(kept),
        }
    except Exception as exc:
        return {
            "status": "degraded",
            "reason": type(exc).__name__[:120],
            "tagged": int(tagged),
            "planned": int(planned),
            "removed": int(removed),
            "kept": int(kept),
        }
    finally:
        if m is not None:
            try:
                m.logout()
            except Exception:
                pass


def safe_summary(payload: dict[str, Any], validation: dict[str, Any]) -> dict[str, Any]:
    return {
        "status": "ready" if validation["ok"] else "blocked",
        "owner_requisites_complete": bool(validation["ok"]),
        "fields_present": validation["fields_present"],
        "fields_required": validation["fields_required"],
        "missing": validation["missing"],
        "invalid_format": validation.get("invalid_format", []),
        "updated_at": validation["updated_at"],
        "card_fingerprint": requisites_fingerprint(payload),
        "org_name_present": bool(str(payload.get("org_name") or "").strip()),
        "banking_present": all(bool(str(payload.get(k) or "").strip()) for k in (
            "bank_account", "bank_bik", "bank_name", "corr_account"
        )),
        "truth": "No banking values are printed. Sending requires --apply and --confirm-share-banking.",
    }


def message_content() -> tuple[str, str]:
    subject = "Re: RE: Партнерская программа"
    body = (
        "Добрый день!\n\n"
        "Направляю актуальную карточку ИП для подготовки клиентского и агентского оформления.\n"
        "Минимальную тестовую конфигурацию подтверждаем: мобильный номер + SIP-транк на 1 канал.\n\n"
        "После оформления пришлите, пожалуйста, SIP-параметры и номер для подключения к нашему Asterisk.\n\n"
        "Спасибо!\nКирилл\nBORIS AI"
    )
    return subject, body



def _mail_dt(value: str | None):
    raw = str(value or "").strip()
    if not raw:
        return None
    try:
        dt = parsedate_to_datetime(raw)
    except Exception:
        return None
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def sent_company_card_copy_exists(
    mailbox_id: int,
    reply_id: str,
    card_fingerprint: str,
    request_date: str = "",
) -> dict[str, Any]:
    """Detect a company card already sent to MCN.

    Prefer BORIS private headers when the provider preserves them. Mail.ru can
    strip those headers when a prepared draft is sent manually, so fall back to
    the read-only Sent-folder attachment proof used by the guardian, bounded by
    the latest provider request date.
    """
    row = _row(mailbox_id)
    if not row:
        return {"exists": False, "reason": "mailbox_not_found"}

    secret = decrypt_secret(row["secret_encrypted"])
    cls = imaplib.IMAP4_SSL if row["imap_ssl"] else imaplib.IMAP4
    m = None
    try:
        m = cls(row["imap_host"], int(row["imap_port"]), timeout=15)
        m.login(row["username"], secret)
        folder = _sent_folder(m)
        if folder:
            st, _ = m.select(folder, readonly=True)
            if st == "OK":
                st, data = m.uid(
                    "search", None,
                    "HEADER", "X-BORIS-Action-ID", f'"{DRAFT_ACTION_ID}"',
                )
                ids = data[0].split() if st == "OK" and data and data[0] else []
                for uid in reversed(ids[-100:]):
                    st, fetched = m.uid(
                        "fetch", uid,
                        "(BODY.PEEK[HEADER.FIELDS (IN-REPLY-TO X-BORIS-CARD-FINGERPRINT)])",
                    )
                    raw = b"".join(
                        x[1] for x in fetched
                        if isinstance(x, tuple) and len(x) > 1 and isinstance(x[1], (bytes, bytearray))
                    ) if st == "OK" else b""
                    if not raw:
                        continue
                    msg = email.message_from_bytes(raw)
                    sent_reply = str(msg.get("In-Reply-To") or "").strip()
                    sent_fp = str(msg.get("X-BORIS-Card-Fingerprint") or "").strip().lower()
                    if str(reply_id or "").strip() and sent_reply != str(reply_id).strip():
                        continue
                    if (
                        card_fingerprint
                        and len(sent_fp) == 64
                        and hmac.compare_digest(sent_fp, str(card_fingerprint).strip().lower())
                    ):
                        return {"exists": True, "reason": "matching_boris_sent_copy"}
    except Exception:
        # Fall through to the independent read-only attachment proof below.
        pass
    finally:
        if m is not None:
            try:
                m.logout()
            except Exception:
                pass

    request_dt = _mail_dt(request_date)
    if request_dt is None:
        return {"exists": False, "reason": "matching_sent_copy_not_found"}

    discovery = ROOT / "scripts" / "phone-company-card-mailbox-discovery.py"
    python_bin = ROOT / "venv" / "bin" / "python"
    try:
        cp = subprocess.run(
            [
                str(python_bin), str(discovery),
                "--mailbox-id", str(int(mailbox_id)),
                "--provider-domain", "mcn.ru",
                "--limit", "250",
            ],
            cwd=str(ROOT),
            text=True,
            capture_output=True,
            timeout=30,
            check=False,
        )
        result = json.loads((cp.stdout or "").strip() or "{}")
    except Exception as exc:
        return {
            "exists": False,
            "reason": ("sent_attachment_scan_error:" + type(exc).__name__)[:120],
        }

    if result.get("status") != "ok":
        return {
            "exists": False,
            "reason": str(result.get("reason") or "sent_attachment_scan_blocked")[:120],
        }

    exact_reply = str(reply_id or "").strip()
    for item in result.get("items") or []:
        if not isinstance(item, dict) or not item.get("company_card_signal"):
            continue
        item_dt = _mail_dt(str(item.get("date") or ""))
        sent_reply = str(item.get("in_reply_to") or "").strip()
        reply_ok = (not exact_reply) or sent_reply == exact_reply
        if item_dt is not None and item_dt >= request_dt and reply_ok:
            return {
                "exists": True,
                "reason": "attachment_seen_after_request",
                "sent_at": item_dt.isoformat(),
            }
    return {"exists": False, "reason": "matching_sent_copy_not_found"}


def prepare_company_card_draft_once(
    mailbox_id: int,
    recipient: str,
    headers: dict[str, str],
    subject: str,
    body: str,
    attachments: list[dict[str, Any]],
    summary: dict[str, Any],
    card_fingerprint: str,
    request_date: str = "",
) -> int:
    """Serialize draft creation across guardian/API/diagnostic processes.

    The lock covers the entire external IMAP boundary, not only the DB state
    write. A concurrent process waits, then re-checks the persisted Message-ID
    and exits without appending a second draft.
    """
    reply_id = str(headers.get("In-Reply-To") or "").strip()
    lock_key = f"mcn-company-card-draft|{int(mailbox_id)}|{reply_id}"
    lock_db = SessionLocal()
    locked = False
    try:
        lock_db.execute(
            text("SELECT pg_advisory_lock(hashtextextended(:k,0))"),
            {"k": lock_key},
        )
        locked = True

        send_state = send_state_for_thread(mailbox_id, reply_id)
        if send_state.get("matches") and send_state.get("status") == "accepted":
            summary.update({
                "status": "already_sent",
                "reason": "exactly_once_guard",
                "message_id": send_state.get("message_id"),
                "retry_blocked": True,
            })
            print(json.dumps(summary, ensure_ascii=False))
            return 0
        if send_state.get("matches") and send_state.get("status") in {"sending", "ambiguous"}:
            summary.update({
                "status": "delivery_ambiguous",
                "reason": "previous_send_may_have_reached_smtp",
                "retry_blocked": True,
            })
            print(json.dumps(summary, ensure_ascii=False))
            return 8

        sent_copy = sent_company_card_copy_exists(mailbox_id, reply_id, card_fingerprint, request_date)
        if sent_copy.get("exists"):
            cleanup = remove_current_company_card_draft(mailbox_id)
            summary.update({
                "status": "already_sent",
                "reason": "sent_folder_exactly_once_guard",
                "retry_blocked": True,
                "draft_cleanup_removed": bool(cleanup.get("removed")),
                "draft_cleanup_reason": str(cleanup.get("reason") or "")[:120] or None,
            })
            print(json.dumps(summary, ensure_ascii=False))
            return 0

        draft_state = company_card_draft_exists(mailbox_id, reply_id)
        if draft_state.get("exists"):
            existing_fingerprint = str(draft_state.get("card_fingerprint") or "").strip()
            if existing_fingerprint and hmac.compare_digest(existing_fingerprint, card_fingerprint):
                summary.update({
                    "status": "draft_exists",
                    "reason": "ok",
                    "draft_fresh": True,
                })
                print(json.dumps(summary, ensure_ascii=False))
                return 0

            stale_cleanup = remove_current_company_card_draft(mailbox_id)
            if draft_state.get("source") == "persisted_message_id" and not stale_cleanup.get("removed"):
                summary.update({
                    "status": "draft_refresh_blocked",
                    "reason": str(stale_cleanup.get("reason") or "stale_draft_cleanup_failed")[:120],
                    "draft_fresh": False,
                })
                print(json.dumps(summary, ensure_ascii=False))
                return 9
            summary["stale_draft_removed"] = bool(stale_cleanup.get("removed"))
            summary["stale_draft_cleanup_reason"] = str(stale_cleanup.get("reason") or "")[:120] or None

        ok, reason, message_id = save_draft(
            mailbox_id,
            recipient,
            subject,
            body,
            headers=headers,
            attachments=attachments,
        )
        state_persisted = False
        if ok and message_id:
            try:
                persist_draft_state(
                    mailbox_id,
                    message_id,
                    reply_id,
                    card_fingerprint,
                )
                state_persisted = True
            except Exception:
                state_persisted = False

        legacy_cleanup = {"removed": False, "reason": "not_attempted", "count": 0}
        if ok and state_persisted:
            legacy_cleanup = remove_legacy_company_card_draft(mailbox_id)
        summary.update({
            "status": "draft_saved" if ok else "draft_failed",
            "reason": reason,
            "message_id": message_id if ok else None,
            "idempotency_state_persisted": state_persisted if ok else False,
            "legacy_draft_removed": bool(legacy_cleanup.get("removed")) if ok else False,
            "legacy_draft_cleanup_reason": str(legacy_cleanup.get("reason") or "")[:120] if ok else None,
        })
        print(json.dumps(summary, ensure_ascii=False))
        return 0 if ok else 7
    finally:
        if locked:
            try:
                lock_db.execute(
                    text("SELECT pg_advisory_unlock(hashtextextended(:k,0))"),
                    {"k": lock_key},
                )
                lock_db.commit()
            except Exception:
                lock_db.rollback()
        lock_db.close()


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--mailbox-id", type=int, default=DEFAULT_MAILBOX_ID)
    ap.add_argument("--recipient", default=DEFAULT_RECIPIENT)
    ap.add_argument("--output", default="")
    ap.add_argument("--request-message-id", default="")
    ap.add_argument("--request-date", default="")
    mode = ap.add_mutually_exclusive_group()
    mode.add_argument("--save-draft", action="store_true")
    mode.add_argument("--apply", action="store_true")
    ap.add_argument("--confirm-share-banking", action="store_true")
    args = ap.parse_args()

    payload = load_owner_requisites()
    validation = validate_requisites(payload)
    summary = safe_summary(payload, validation)
    if not validation["ok"]:
        print(json.dumps(summary, ensure_ascii=False))
        return 3

    recipient_check = validate_mcn_recipient(args.recipient)
    if not recipient_check["ok"]:
        summary.update({
            "status": "blocked",
            "reason": "recipient_not_mcn_domain",
        })
        print(json.dumps(summary, ensure_ascii=False))
        return 6
    recipient = str(recipient_check["address"])
    summary["recipient_domain_verified"] = True

    if args.apply and not args.confirm_share_banking:
        summary.update({
            "status": "blocked",
            "reason": "banking_share_confirmation_required",
        })
        print(json.dumps(summary, ensure_ascii=False))
        return 4

    temp_dir = None
    if args.output:
        card_path = Path(args.output).expanduser().resolve()
    else:
        temp_dir = tempfile.TemporaryDirectory(prefix="boris-mcn-card-")
        card_path = Path(temp_dir.name) / "Карточка_ИП_актуальная.docx"

    try:
        build_company_card(payload, card_path)
        summary["document_built"] = True
        summary["document_size"] = card_path.stat().st_size

        if not args.apply and not args.save_draft:
            summary["status"] = "ready_to_send"
            if args.output:
                summary["output"] = str(card_path)
            print(json.dumps(summary, ensure_ascii=False))
            return 0

        headers = latest_mcn_thread_headers(
            args.mailbox_id,
            request_message_id=str(args.request_message_id or "").strip(),
            request_date=str(args.request_date or "").strip(),
        )
        provider_request_date = str(headers.pop("_request_date", "") or "").strip()
        card_fingerprint = requisites_fingerprint(payload)
        headers["X-BORIS-Action-ID"] = DRAFT_ACTION_ID
        headers["X-BORIS-Card-Fingerprint"] = card_fingerprint
        subject, body = message_content()
        attachments = [{
            "path": str(card_path),
            "filename": card_path.name,
            "mime": DOCX_MIME,
        }]

        if args.save_draft:
            return prepare_company_card_draft_once(
                args.mailbox_id,
                recipient,
                headers,
                subject,
                body,
                attachments,
                summary,
                card_fingerprint,
                provider_request_date,
            )

        reply_id = str(headers.get("In-Reply-To") or "").strip()
        draft_state = load_draft_state(DRAFT_STATE_KEY)
        prepared_fingerprint = str(draft_state.get("card_fingerprint") or "")
        prepared_reply = str(draft_state.get("in_reply_to") or "")
        prepared_action = str(draft_state.get("action_id") or "")
        if (
            prepared_action != DRAFT_ACTION_ID
            or prepared_reply != reply_id
            or prepared_fingerprint != card_fingerprint
        ):
            summary.update({
                "status": "blocked",
                "reason": "prepared_card_changed_or_stale",
                "retry_blocked": True,
            })
            print(json.dumps(summary, ensure_ascii=False))
            return 9
        draft_state_live = company_card_draft_exists(args.mailbox_id, reply_id)
        if not draft_state_live.get("exists"):
            summary.update({
                "status": "blocked",
                "reason": "prepared_draft_missing",
                "retry_blocked": True,
            })
            print(json.dumps(summary, ensure_ascii=False))
            return 9
        claim = claim_send_once(args.mailbox_id, reply_id)
        if not claim.get("claimed"):
            previous = str(claim.get("status") or "")
            if previous == "accepted":
                previous_mid = str(claim.get("message_id") or "")
                summary.update({
                    "status": "already_sent",
                    "reason": "exactly_once_guard",
                    "message_id": previous_mid or None,
                    "sent_copy_saved": bool(
                        sent_copy_saved(args.mailbox_id, previous_mid)
                    ) if previous_mid else False,
                })
                print(json.dumps(summary, ensure_ascii=False))
                return 0
            summary.update({
                "status": "delivery_ambiguous",
                "reason": (
                    "previous_send_may_have_reached_smtp"
                    if previous in {"sending", "ambiguous"}
                    else "send_claim_unavailable"
                ),
                "retry_blocked": True,
            })
            print(json.dumps(summary, ensure_ascii=False))
            return 8

        ok, reason, message_id = send_outbound(
            args.mailbox_id,
            recipient,
            subject,
            body,
            headers=headers,
            attachments=attachments,
        )
        if ok:
            state_persisted = True
            try:
                finish_send_state(
                    args.mailbox_id,
                    reply_id,
                    status="accepted",
                    message_id=message_id,
                    reason=reason,
                )
            except Exception:
                # The durable pre-send claim remains "sending", which blocks a
                # duplicate retry even if this final bookkeeping write failed.
                state_persisted = False
            copy_saved = bool(
                sent_copy_saved(args.mailbox_id, message_id)
            ) if message_id else False
            draft_cleanup = remove_current_company_card_draft(args.mailbox_id)
            summary.update({
                "status": "sent" if state_persisted else "sent_state_pending",
                "reason": reason,
                "message_id": message_id or None,
                "sent_copy_saved": copy_saved,
                "send_state_persisted": state_persisted,
                "draft_removed_after_send": bool(draft_cleanup.get("removed")),
                "draft_cleanup_reason": str(draft_cleanup.get("reason") or "")[:120],
                "retry_blocked": True,
            })
            print(json.dumps(summary, ensure_ascii=False))
            # SMTP acceptance is the external delivery boundary. Missing Sent
            # copy is repaired by mailbox_sent_copy_queue and must never cause
            # the business message to be transmitted twice.
            return 0

        delivery_unknown = str(reason or "").startswith("delivery_unknown:")
        final_state = "ambiguous" if delivery_unknown else "safe_failure"
        state_persisted = True
        try:
            finish_send_state(
                args.mailbox_id,
                reply_id,
                status=final_state,
                message_id=message_id,
                reason=reason,
            )
        except Exception:
            state_persisted = False
        summary.update({
            "status": "delivery_ambiguous" if delivery_unknown or not state_persisted else "send_failed",
            "reason": reason,
            "message_id": None,
            "send_state_persisted": state_persisted,
            "retry_blocked": bool(delivery_unknown or not state_persisted),
        })
        print(json.dumps(summary, ensure_ascii=False))
        return 8 if delivery_unknown or not state_persisted else 5
    finally:
        if temp_dir is not None:
            temp_dir.cleanup()


if __name__ == "__main__":
    raise SystemExit(main())
