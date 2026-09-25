#!/usr/bin/env python3
"""Read-only discovery of company-card attachments in a BORIS mailbox Sent folder.

Safety contract:
- IMAP Sent folder is selected readonly.
- BODY.PEEK is used, so no message flags/cursors are changed.
- Mailbox passwords and attachment bytes/content are never printed.
- Only messages addressed to the provider domain are inspected.
- Output is limited to message metadata and attachment filenames.
"""
from __future__ import annotations

import argparse
import email
import io
import imaplib
import importlib.util
import json
import re
import sys
from email.utils import getaddresses
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.services.client_mailboxes import _row, _sent_folder

_DISCOVERY = ROOT / "scripts" / "mcn-mailbox-discovery.py"
_spec = importlib.util.spec_from_file_location("mcn_mailbox_discovery_safe", _DISCOVERY)
_D = importlib.util.module_from_spec(_spec)
assert _spec and _spec.loader
_spec.loader.exec_module(_D)

_CARD_SCRIPT = ROOT / "scripts" / "phone-mcn-company-card-send.py"
_card_spec = importlib.util.spec_from_file_location("mcn_company_card_current_safe", _CARD_SCRIPT)
_C = importlib.util.module_from_spec(_card_spec)
assert _card_spec and _card_spec.loader
_card_spec.loader.exec_module(_C)

CARD_RE = re.compile(
    r"(карточ|реквиз|company.?card|details|инн|огрн|банк|bank|ип\b)",
    re.I,
)


def _decode(value: str | None) -> str:
    try:
        from email.header import decode_header, make_header
        return str(make_header(decode_header(value or "")))
    except Exception:
        return str(value or "")


def _provider_recipient(msg: email.message.Message, domain: str) -> bool:
    values = []
    for header in ("To", "Cc", "Bcc"):
        values.extend(msg.get_all(header, []))
    target = domain.lower().lstrip("@")
    for _, addr in getaddresses(values):
        d = addr.rsplit("@", 1)[-1].lower() if "@" in addr else ""
        if d == target or d.endswith("." + target):
            return True
    return False



def _norm(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def _current_card_payload() -> dict | None:
    try:
        payload = _C.load_owner_requisites()
        check = _C.validate_requisites(payload)
        return payload if check.get("ok") else None
    except Exception:
        return None


def _docx_matches_current_card(blob: bytes, current: dict | None) -> bool | None:
    if not blob or not current:
        return None
    try:
        doc = Document(io.BytesIO(blob))
        parts = [p.text for p in doc.paragraphs]
        parts.extend(
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        )
        text_blob = _norm("\n".join(parts))
    except Exception:
        return False

    for key in _C.REQUIRED:
        value = _norm(current.get(key))
        if not value or value not in text_blob:
            return False
    return True


def discover(mailbox_id: int, provider_domain: str = "mcn.ru", limit: int = 1200) -> dict:
    row = _row(int(mailbox_id))
    if not row:
        return {"status": "blocked", "reason": "mailbox_not_found", "mailbox_id": mailbox_id}
    # Standalone diagnostics do not inherit systemd's EnvironmentFile. Reuse
    # the canonical MCN scanner's file-backed Fernet reader so an encrypted
    # mailbox password is never accidentally sent verbatim to IMAP.
    secret = _D._decrypt_mailbox_secret(row["secret_encrypted"])
    if not secret:
        return {"status": "blocked", "reason": "mailbox_secret_unavailable", "mailbox_id": mailbox_id}

    cls = imaplib.IMAP4_SSL if row["imap_ssl"] else imaplib.IMAP4
    conn = None
    try:
        conn = cls(row["imap_host"], int(row["imap_port"]), timeout=12)
        conn.login(row["username"], secret)
        folder = _sent_folder(conn)
        if not folder:
            return {"status": "blocked", "reason": "sent_folder_not_found", "mailbox_id": mailbox_id}
        status, count_data = conn.select(folder, readonly=True)
        if status != "OK":
            return {"status": "blocked", "reason": "sent_select_failed", "mailbox_id": mailbox_id}

        # Server-side recipient narrowing first. If the provider does not support
        # partial-domain TO search, bounded ALL fallback is still readonly.
        status, data = conn.uid("search", None, "TO", f'"{provider_domain}"')
        uids = data[0].split() if status == "OK" and data and data[0] else []
        if not uids:
            status, data = conn.uid("search", None, "ALL")
            uids = data[0].split() if status == "OK" and data and data[0] else []
        chosen = uids[-max(1, min(int(limit), 5000)):]

        current_card = _current_card_payload()
        items = []
        scanned = 0
        for uid_raw in reversed(chosen):
            status, fetched = conn.uid("fetch", uid_raw, "(BODY.PEEK[])")
            if status != "OK" or not fetched:
                continue
            raw = b"".join(
                item[1]
                for item in fetched
                if isinstance(item, tuple) and len(item) > 1 and isinstance(item[1], (bytes, bytearray))
            )
            if not raw:
                continue
            scanned += 1
            msg = email.message_from_bytes(raw)
            if not _provider_recipient(msg, provider_domain):
                continue
            attachments = []
            current_card_match = None
            for part in msg.walk():
                filename = part.get_filename()
                if filename:
                    decoded_name = _decode(filename)[:300]
                    attachments.append(decoded_name)
                    if CARD_RE.search(decoded_name) and decoded_name.lower().endswith(".docx"):
                        match = _docx_matches_current_card(part.get_payload(decode=True) or b"", current_card)
                        if match is True:
                            current_card_match = True
                        elif match is False and current_card_match is None:
                            current_card_match = False
            subject = _decode(msg.get("Subject"))[:500]
            names_blob = " ".join(attachments)
            action_id = _decode(msg.get("X-BORIS-Action-ID"))[:160]
            in_reply_to = _decode(msg.get("In-Reply-To"))[:500]
            message_id = _decode(msg.get("Message-ID"))[:500]
            company_card_signal = bool(CARD_RE.search(names_blob))
            items.append({
                "uid": uid_raw.decode(errors="ignore"),
                "date": _decode(msg.get("Date"))[:200],
                "subject": subject,
                "message_id": message_id,
                "in_reply_to": in_reply_to,
                "action_id": action_id,
                "attachments": attachments,
                "company_card_signal": company_card_signal,
                "current_card_match": current_card_match,
                "boris_action_signal": bool(
                    company_card_signal
                    and action_id.startswith("boris-mcn-company-card-v")
                ),
            })
            if len(items) >= 100:
                break
        return {
            "status": "ok",
            "mailbox_id": int(mailbox_id),
            "provider_domain": provider_domain,
            "sent_folder": folder,
            "messages_available": int(count_data[0]) if count_data and count_data[0] else None,
            "messages_scanned": scanned,
            "matched_messages": len(items),
            "company_card_candidates": sum(1 for x in items if x["company_card_signal"]),
            "items": items,
            "truth": "read-only Sent discovery; attachment bytes/content are used only in-memory for boolean current-card verification and are never returned",
        }
    except Exception as exc:
        return {
            "status": "blocked",
            "reason": f"{type(exc).__name__}:{' '.join(str(exc).split())[:160]}",
            "mailbox_id": int(mailbox_id),
        }
    finally:
        if conn is not None:
            try:
                conn.logout()
            except Exception:
                pass


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--mailbox-id", type=int, default=2)
    ap.add_argument("--provider-domain", default="mcn.ru")
    ap.add_argument("--limit", type=int, default=1200)
    args = ap.parse_args()
    out = discover(args.mailbox_id, args.provider_domain, args.limit)
    print(json.dumps(out, ensure_ascii=False, indent=2, default=str))
    return 0 if out.get("status") == "ok" else 2


if __name__ == "__main__":
    raise SystemExit(main())
